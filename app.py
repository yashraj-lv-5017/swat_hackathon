import streamlit as st
import pandas as pd
import plotly.express as px
from sentinel import run_sentinel_audit
from fpdf import FPDF
from fpdf.enums import XPos, YPos
import diff_match_patch as dmp_module
import pdfplumber
from docx import Document
import time

# --- 1. PAGE CONFIGURATION ---
st.set_page_config(
    page_title="Axion Sentinel Elite: Governance Engine",
    layout="wide",
    page_icon="🛡️"
)

# --- 2. SESSION STATE ---
if 'audit_history' not in st.session_state:
    st.session_state.audit_history = []
if 'current_text' not in st.session_state:
    st.session_state.current_text = ""

# --- 3. UTILITIES ---
def analyze_complexity(text):
    words = text.split()
    if not words: return 0
    return round(sum(len(w) for w in words) / len(words), 2)

def speak_text(text):
    """Injects JavaScript to handle Text-to-Speech in the browser."""
    safe_text = text.replace("'", "\\'").replace("\n", " ")
    js_code = f"""
    <script>
        var msg = new SpeechSynthesisUtterance();
        msg.text = '{safe_text}';
        msg.rate = 0.95; 
        window.speechSynthesis.speak(msg);
    </script>
    """
    st.components.v1.html(js_code, height=0)

def extract_text_from_file(uploaded_file):
    text = ""
    try:
        if uploaded_file.type == "application/pdf":
            with pdfplumber.open(uploaded_file) as pdf:
                text = "\n".join([page.extract_text() for page in pdf.pages if page.extract_text()])
        elif uploaded_file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            doc = Document(uploaded_file)
            text = "\n".join([para.text for para in doc.paragraphs])
        elif uploaded_file.type in ["text/plain"]:
            text = uploaded_file.read().decode("utf-8")
    except Exception as e:
        st.error(f"Error reading file: {e}")
    return text

def render_diff(old_text, new_text):
    dmp = dmp_module.diff_match_patch()
    diffs = dmp.diff_main(old_text, new_text)
    dmp.diff_cleanupSemantic(diffs)
    html_result = ""
    for (flag, data) in diffs:
        if flag == 0: html_result += f"<span>{data}</span>"
        elif flag == 1: html_result += f"<span style='background-color: #d4edda; color: #155724; font-weight: bold;'>{data}</span>"
        elif flag == -1: html_result += f"<span style='background-color: #f8d7da; color: #721c24; text-decoration: line-through;'>{data}</span>"
    st.markdown(f"<div style='border: 1px solid #ddd; padding: 20px; border-radius: 10px; line-height: 1.8; background-color: white; color: black;'>{html_result}</div>", unsafe_allow_html=True)

def create_pdf_report(report, audience, complexity_delta):
    pdf = FPDF()
    pdf.add_page()
    def sanitize(text):
        if not text: return ""
        rep = {"’": "'", "‘": "'", "“": '"', "”": '"', "—": "-", "–": "-"}
        for k, v in rep.items(): text = text.replace(k, v)
        return str(text).encode("latin-1", "replace").decode("latin-1")

    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, sanitize(f"Axion Governance Audit: {audience}"), new_x=XPos.LMARGIN, new_y=YPos.NEXT, align='C')
    pdf.ln(10)
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 10, sanitize(f"Brand Health Score: {report.get('brand_health_score', 0)}%"), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.cell(0, 10, sanitize(f"Linguistic Shift: {complexity_delta}"), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(10)
    pdf.set_font("Helvetica", "B", 11)
    pdf.cell(0, 10, "Approved Content:", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Helvetica", "", 10)
    pdf.multi_cell(0, 8, sanitize(report.get('adapted_text', "")))
    return bytes(pdf.output())

# --- 4. SIDEBAR ---
with st.sidebar:
    st.header("🎯 Context Matrix")
    channel = st.selectbox("Distribution Channel", ["LinkedIn post", "Marketing email", "Press Release"])
    audience = st.selectbox("Target Stakeholder", ["Executive", "Practitioner", "Technical Lead"])
    st.divider()
    st.subheader("📋 Session History")
    if st.session_state.audit_history:
        for entry in reversed(st.session_state.audit_history):
            st.caption(f"{entry['time']} | Score: {entry['score']}% ({entry['audience']})")
    st.divider()
    show_debug = st.toggle("🔍 Trace Mode (RAG/Agent Logic)")

# --- 5. MAIN CONTENT ---
st.title("🛡️ Axion Sentinel: Enterprise Governance")
st.subheader("📁 Content Ingestion")
uploaded_file = st.file_uploader("Upload a Draft", type=["pdf", "docx", "txt"])
if uploaded_file:
    st.session_state.current_text = extract_text_from_file(uploaded_file)

source_text = st.text_area("Input Draft Content:", value=st.session_state.current_text, height=180)

# Container for layout control
results_container = st.container()

# --- 6. AUDIT EXECUTION ---
if st.button("🚀 Execute Strategic Audit", use_container_width=True):
    if not source_text:
        st.error("Buffer is empty. Please provide text or upload a file.")
    else:
        my_bar = st.progress(0, text="Initializing Audit Engine...")
        
        # CORE CALL - Results stored in Session State to prevent reset on Voice click
        st.session_state.report = run_sentinel_audit(source_text, channel, audience)
        
        orig_comp = analyze_complexity(source_text)
        new_comp = analyze_complexity(st.session_state.report.get('adapted_text', ''))
        st.session_state.comp_delta = round(new_comp - orig_comp, 2)
        
        # Log to history
        st.session_state.audit_history.append({
            "time": time.strftime("%H:%M"), 
            "score": st.session_state.report.get('brand_health_score', 0), 
            "audience": audience
        })
        
        my_bar.progress(100, text="✅ Audit Complete")
        time.sleep(0.5); my_bar.empty()

# --- 7. DISPLAY RESULTS (Persistent View) ---
if 'report' in st.session_state:
    report = st.session_state.report
    comp_delta = st.session_state.comp_delta
    score = report.get('brand_health_score', 0)
    new_comp = analyze_complexity(report.get('adapted_text', ''))

    # TOP METRICS
    m1, m2, m3, m4 = st.columns(4)
    m1.metric("Health Score", f"{score}%", f"{score-100}%", delta_color="inverse")
    m2.metric("Audit Confidence", f"{report.get('confidence_score', 65)}%")
    m3.metric("Linguistic Shift", f"{new_comp}", delta=comp_delta)
    m4.metric("Violations", len(report.get('violations', [])))

    # --- NEW: BUSINESS IMPACT ANALYTICS (ROI) ---
    st.divider()
    st.subheader("💼 Enterprise Business Impact")
    
    audit_count = len(st.session_state.audit_history)
    time_saved = audit_count * 0.75  # Assume 45 mins saved per manual audit
    dollars_saved = time_saved * 80   # Assume $80/hr labor cost
    
    # Market Readiness Logic
    target_comp = 5.5 if audience == "Executive" else 8.0
    market_score = max(0, 100 - (abs(new_comp - target_comp) * 15))

    bi1, bi2, bi3 = st.columns(3)
    with bi1:
        st.metric("Governance Savings", f"${dollars_saved:,.0f}", help="Manual brand review labor replacement.")
    with bi2:
        st.metric("Efficiency Gain", f"{time_saved:.1f} hrs", delta="Real-time")
    with bi3:
        st.metric("Market Readiness", f"{int(market_score)}%", help="Persona-channel alignment score.")

    if audit_count > 1:
        history_df = pd.DataFrame(st.session_state.audit_history)
        st.line_chart(history_df['score'], height=150)
        st.caption("Historical Brand Health Trend")

    # ORIGINAL CHARTS
    if report.get('violations'):
        st.subheader("📊 Strategic Violation Analysis")
        df = pd.DataFrame(report['violations'])
        if 'category' not in df.columns: df['category'] = 'General'
        c1, c2 = st.columns(2)
        with c1:
            st.plotly_chart(px.pie(df, names='category', title="Violation Mix", hole=0.5, 
                                   color_discrete_sequence=px.colors.qualitative.Pastel), use_container_width=True)
        with c2:
            y_val = 'impact_score' if 'impact_score' in df.columns else None
            st.plotly_chart(px.bar(df, x='category', y=y_val, title="Risk Severity Density", 
                                   color='category', color_discrete_sequence=px.colors.qualitative.Safe), use_container_width=True)

    # --- TRANSFORMATION & VOICE ASSISTANT ---
    st.divider()
    st.subheader(f"🔄 Transformation Insight: {audience} Persona")
    
    col_view, col_voice = st.columns([4, 1])
    with col_view:
        render_diff(source_text, report.get('adapted_text', ''))
    with col_voice:
        st.write("") # UI Padding
        if st.button("🔊 Play Voice"):
            speak_text(report.get('adapted_text', ''))
            st.success("Playing Audio...")

    # CHANGE LOG & PDF EXPORT
    st.subheader("📝 Strategic Change Log")
    st.table(report.get('change_log', []))
    
    pdf_bytes = create_pdf_report(report, audience, comp_delta)
    st.download_button("📥 Download Official Audit Report", data=pdf_bytes, file_name=f"Axion_Audit_{audience}.pdf", use_container_width=True)

    if show_debug:
        with st.expander("🔍 JSON System Trace"):
            st.json(report)